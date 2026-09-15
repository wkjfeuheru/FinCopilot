"""Markdown -> docx 导出（文档 3.9.3）。

完全基于 python-docx 从零构建，使样式成为代码（因而可测试），
而不是一个无法在 diff 中审阅的二进制模板。
"""

from __future__ import annotations

import re
from pathlib import Path

HEADING_RE = re.compile(r"^(#{1,4})\s+(.*)$")
IMAGE_RE = re.compile(r"^!\[(.*?)\]\((.+?)\)\s*$")
TABLE_ROW_RE = re.compile(r"^\|.*\|\s*$")
TABLE_SEP_RE = re.compile(r"^\|[\s:|-]+\|\s*$")

BODY_FONT = "SimSun"
HEADING_FONT = "SimHei"
MAX_IMAGE_INCHES = 6.0

DISCLAIMER = "本报告由 FinHarness 依据公开数据自动生成，仅供研究参考，不构成投资建议。"


class DocxExportError(RuntimeError):
    """无法生成文档时抛出。"""


def _set_run_font(run, *, font: str) -> None:
    """同时为 ASCII 与东亚文字两个字体槽位设置字体。

    python-docx 默认只设置 ASCII 字体；若缺少 ``w:eastAsia``，Word 会回退到
    拉丁字体，导致中文字形渲染不一致。
    """
    from docx.oxml.ns import qn

    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)


def export_markdown_to_docx(markdown: str, *, out_path: str | Path, topic: str = "") -> Path:
    """把报告 markdown 转换为 ``.docx``；返回写入的路径。"""
    try:
        from docx import Document
        from docx.shared import Inches
    except ImportError as exc:  # pragma: no cover - python-docx 为必需依赖
        raise DocxExportError(f"python-docx 不可用：{exc}") from exc

    path = Path(out_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()

    lines = markdown.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].rstrip()

        if not line.strip():
            index += 1
            continue

        # 表格：一个表头行、一个分隔行，随后是正文行。
        if TABLE_ROW_RE.match(line) and index + 1 < len(lines) and TABLE_SEP_RE.match(lines[index + 1]):
            block, index = _consume_table(lines, index)
            _write_table(document, block, font=BODY_FONT)
            continue

        image = IMAGE_RE.match(line)
        if image:
            caption, target = image.group(1), image.group(2).strip()
            # 含空格的路径会用尖括号包裹以满足 Markdown 语法；
            # 尖括号是语法，不属于文件系统路径本身。
            if target.startswith("<") and target.endswith(">"):
                target = target[1:-1].replace("\\>", ">").strip()
            # 仅嵌入本地文件；缺失的图片降级为一条提示，
            # 使文档其余部分仍能导出。
            candidate = Path(target)
            if candidate.is_file():
                try:
                    document.add_picture(str(candidate), width=Inches(MAX_IMAGE_INCHES))
                except Exception as exc:  # noqa: BLE001 - 坏图片不得中断导出
                    _add_note(document, f"[图片嵌入失败：{target}（{exc}）]", font=BODY_FONT)
                else:
                    if caption:
                        _add_caption(document, caption, font=BODY_FONT)
            else:
                _add_note(document, f"[图片缺失：{target}]", font=BODY_FONT)
            index += 1
            continue

        heading = HEADING_RE.match(line)
        if heading:
            level = len(heading.group(1))
            text = _strip_inline(heading.group(2))
            paragraph = document.add_heading(level=level)
            run = paragraph.add_run(text)
            _set_run_font(run, font=HEADING_FONT)
            index += 1
            continue

        paragraph = document.add_paragraph()
        run = paragraph.add_run(_strip_inline(line))
        _set_run_font(run, font=BODY_FONT)
        index += 1

    document.save(str(path))
    return path


def _consume_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    """收集从 ``start`` 起连续的表格行，返回单元格二维列表与下一行索引。"""
    block: list[list[str]] = []
    index = start
    while index < len(lines) and TABLE_ROW_RE.match(lines[index]):
        if TABLE_SEP_RE.match(lines[index]):
            index += 1
            continue
        cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        block.append(cells)
        index += 1
    return block, index


def _write_table(document, block: list[list[str]], *, font: str) -> None:
    """把单元格二维列表写入 docx 表格，并对每行文本套用指定字体。"""
    if not block:
        return
    width = max(len(row) for row in block)
    table = document.add_table(rows=0, cols=width)
    table.style = "Table Grid"
    for row in block:
        cells = table.add_row().cells
        for position in range(width):
            text = row[position] if position < len(row) else ""
            paragraph = cells[position].paragraphs[0]
            run = paragraph.add_run(_strip_inline(text))
            _set_run_font(run, font=font)


def _add_note(document, text: str, *, font: str) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.italic = True
    _set_run_font(run, font=font)


def _add_caption(document, text: str, *, font: str) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.font.size = None  # 继承正文大小
    _set_run_font(run, font=font)


def _strip_inline(text: str) -> str:
    """去掉 docx 不需要的 markdown 强调标记。"""
    return text.replace("**", "").replace("`", "").strip()
